from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .ir import Block, DocIR

HEADING_SCALE = {1: 1.6, 2: 1.35, 3: 1.15}
SKIP_TYPES = {"pageNumber", "figure", "formula", "table", "header", "footer"}


def _set_east_asian(element: Any, font_name: str) -> None:
    rpr = element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)


def _font_name() -> str:
    return os.environ.get("PDF_FONT_NAME", "Noto Sans CJK SC")


def _setup_document(ir: DocIR, base: float) -> Document:
    document = Document()
    if ir.pages:
        section = document.sections[0]
        section.page_width = Pt(ir.pages[0].width)
        section.page_height = Pt(ir.pages[0].height)
        section.left_margin = Cm(2.0)
        section.right_margin = Cm(2.0)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.8)

    font = _font_name()
    normal = document.styles["Normal"]
    normal.font.name = font
    normal.font.size = Pt(base)
    _set_east_asian(normal.element, font)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.5

    if ir.title:
        document.core_properties.title = ir.title
    return document


def _add_heading(document: Document, block: Block, text: str, base: float) -> None:
    level = min(3, max(1, int(block.level or 1)))
    paragraph = document.add_paragraph(style="Heading " + str(level))
    run = paragraph.add_run(text)
    run.bold = True
    run.font.size = Pt(base * HEADING_SCALE[level])
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = _font_name()
    _set_east_asian(run._element, _font_name())
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.2


def _add_caption(document: Document, text: str, base: float) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(base * 0.9)
    run.font.name = _font_name()
    _set_east_asian(run._element, _font_name())
    paragraph.paragraph_format.space_after = Pt(4)


def _add_code(document: Document, text: str, base: float) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(base * 0.9)
    paragraph.paragraph_format.left_indent = Cm(0.4)
    paragraph.paragraph_format.space_after = Pt(4)


def _add_footnote(document: Document, text: str, base: float) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(text)
    run.font.size = Pt(base * 0.85)
    run.font.name = _font_name()
    _set_east_asian(run._element, _font_name())
    paragraph.paragraph_format.space_after = Pt(2)


def render_docx(ir: DocIR, path: Path) -> dict[str, Any]:
    base = min(max(float(ir.body_size or 10.5), 9.0), 12.0)
    document = _setup_document(ir, base)

    rendered = 0
    skipped = 0
    for page in ir.pages:
        for block in sorted(page.blocks, key=lambda item: item.reading_order):
            if block.type in SKIP_TYPES:
                skipped += 1
                continue
            text = block.final_text()
            if not text:
                skipped += 1
                continue
            if block.type == "heading":
                _add_heading(document, block, text, base)
            elif block.type == "caption":
                _add_caption(document, text, base)
            elif block.type == "code":
                _add_code(document, text, base)
            elif block.type == "footnote":
                _add_footnote(document, text, base)
            elif block.type == "listItem":
                paragraph = document.add_paragraph(text)
                paragraph.paragraph_format.left_indent = Cm(0.6)
                paragraph.paragraph_format.space_after = Pt(4)
            else:
                document.add_paragraph(text)
            rendered += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
    return {"rendered": rendered, "skipped": skipped, "base_size": base}
