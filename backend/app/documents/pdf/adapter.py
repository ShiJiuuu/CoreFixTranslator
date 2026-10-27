from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from docx import Document

from ..base import DocumentError, ParsedDocument, Unit
from .parser import parse_pdf
from .render_docx import render_docx
from .render_pdf import convert_docx_to_pdf

try:
    import fitz
except Exception:
    fitz = None

TRANSLATABLE_TYPES = {"heading", "paragraph", "listItem", "caption", "footnote"}


class PdfAdapter:
    extensions = (".pdf",)

    @property
    def output_suffix(self) -> str:
        mode = os.environ.get("PDF_RENDER", "pdf").strip().lower()
        return ".docx" if mode == "docx" else ".pdf"

    def parse(self, path: Path, options: dict[str, Any] | None = None) -> ParsedDocument:
        ir = parse_pdf(path)
        units: list[Unit] = []
        for block in ir.ordered_blocks():
            if block.type not in TRANSLATABLE_TYPES:
                continue
            if not block.source_text.strip():
                continue
            units.append(
                Unit(
                    id=block.id,
                    text=block.source_text,
                    location={
                        "type": block.type,
                        "page": block.page_index,
                        "bbox": block.bbox.as_list(),
                        "level": block.level,
                    },
                )
            )

        meta = {
            "pages": len(ir.pages),
            "blocks": ir.stats.get("blocks", 0),
            "types": ir.stats.get("types", {}),
            "body_size": ir.body_size,
            "translatable": len(units),
            "images": ir.stats.get("images", 0),
            "render_mode": self.output_suffix,
        }
        return ParsedDocument(units=units, meta=meta, handle={"ir": ir})

    def rebuild(self, parsed: ParsedDocument, translations: dict[str, str], out: Path) -> None:
        ir = parsed.handle["ir"]
        for page in ir.pages:
            for block in page.blocks:
                if block.id in translations:
                    block.target_text = translations[block.id]

        if out.suffix.lower() == ".docx":
            render_docx(ir, out)
            return

        temp_docx = out.with_name(out.stem + ".source.docx")
        try:
            render_docx(ir, temp_docx)
            convert_docx_to_pdf(temp_docx, out)
        finally:
            temp_docx.unlink(missing_ok=True)

    def verify(self, source: Path, out: Path, parsed: ParsedDocument) -> list[dict[str, Any]]:
        if out.suffix.lower() == ".docx":
            try:
                document = Document(str(out))
            except Exception as exc:
                return [{"check": "reopen", "ok": False, "detail": str(exc)}]
            return [
                {"check": "paragraphs", "ok": len(document.paragraphs) > 0, "count": len(document.paragraphs)},
            ]

        if fitz is None:
            return [{"check": "pdf_library", "ok": False, "detail": "PyMuPDF missing"}]
        try:
            document = fitz.open(str(out))
        except Exception as exc:
            return [{"check": "reopen", "ok": False, "detail": str(exc)}]
        text_chars = sum(len(page.get_text().strip()) for page in document)
        pages = document.page_count
        document.close()
        return [
            {"check": "pages", "ok": pages > 0, "pages": pages},
            {"check": "has_text", "ok": text_chars > 0, "chars": text_chars},
        ]
