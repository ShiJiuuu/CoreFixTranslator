from __future__ import annotations

import os
import shutil
import statistics
import subprocess
import tempfile
import uuid
from pathlib import Path
from typing import Any, Iterator

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .base import DocumentError, ParsedDocument, Unit

try:
    import fitz
except Exception:
    fitz = None


def _is_cjk(char: str) -> bool:
    return ("\u3400" <= char <= "\u9fff") or ("\uf900" <= char <= "\ufaff") or ("\u3000" <= char <= "\u303f")


def flow_text(text: str) -> str:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if not lines:
        return ""
    result = lines[0]
    for line in lines[1:]:
        if result and _is_cjk(result[-1]) and _is_cjk(line[0]):
            result += line
        else:
            result += " " + line
    return result


def find_soffice() -> str | None:
    configured = os.environ.get("SOFFICE_BIN", "").strip()
    if configured and Path(configured).exists():
        return configured
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    for candidate in (
        "/usr/bin/soffice",
        "/usr/lib/libreoffice/program/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if Path(candidate).exists():
            return candidate
    return None


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path, *, binary: str | None = None, timeout: int | None = None) -> None:
    soffice = binary or find_soffice()
    if not soffice:
        raise DocumentError("pdf_converter_missing", "LibreOffice not found. Install libreoffice-writer or set SOFFICE_BIN.")
    limit = int(timeout or os.environ.get("PDF_CONVERT_TIMEOUT_SECONDS", "300"))
    profile_dir = Path(tempfile.gettempdir()) / ("corefix-lo-" + uuid.uuid4().hex)
    profile_dir.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as outdir:
        command = [
            soffice,
            "--headless",
            "--norestore",
            "--nolockcheck",
            "--nodefault",
            "-env:UserInstallation=file://" + str(profile_dir),
            "--convert-to",
            "pdf",
            "--outdir",
            outdir,
            str(docx_path),
        ]
        try:
            result = subprocess.run(command, capture_output=True, timeout=limit, text=True)
        except FileNotFoundError as exc:
            raise DocumentError("pdf_converter_missing", "LibreOffice binary not executable.") from exc
        except subprocess.TimeoutExpired as exc:
            raise DocumentError("pdf_conversion_timeout", "LibreOffice conversion timed out.") from exc

        produced = Path(outdir) / (docx_path.stem + ".pdf")
        if result.returncode != 0 or not produced.exists():
            detail = (result.stderr or result.stdout or "").strip()[-400:]
            raise DocumentError("pdf_conversion_failed", "LibreOffice failed: " + (detail or "unknown error"))
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(produced), str(pdf_path))

    shutil.rmtree(profile_dir, ignore_errors=True)


def build_docx(parsed: ParsedDocument, translations: dict[str, str], path: Path) -> None:
    document = Document()
    section = document.sections[0]
    width = float(parsed.meta.get("width") or 595.0)
    height = float(parsed.meta.get("height") or 842.0)
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    font_name = os.environ.get("PDF_FONT_NAME", "Noto Sans CJK SC")
    style = document.styles["Normal"]
    style.font.name = font_name
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), font_name)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.15

    for unit in parsed.units:
        text = flow_text(translations.get(unit.id) or unit.text)
        if not text:
            continue
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15

    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))


def _rgb(value: int) -> tuple[float, float, float]:
    return (((value >> 16) & 255) / 255.0, ((value >> 8) & 255) / 255.0, (value & 255) / 255.0)


def _insert_fitted(page: Any, rect: Any, text: str, base_size: float, color: tuple[float, float, float]) -> bool:
    size = max(3.5, min(base_size, 24.0))
    while size >= 3.5:
        try:
            remaining = page.insert_textbox(rect, text, fontname="F0", fontsize=size, color=color, align=0)
        except Exception:
            return False
        if remaining >= 0:
            return True
        size -= 0.5
    try:
        page.insert_text((rect.x0, min(rect.y1, rect.y0 + max(6.0, base_size))), text, fontname="F0", fontsize=3.5, color=color)
        return True
    except Exception:
        return False


def _dominant_size(sizes: list[float], default: float = 10.0) -> float:
    if not sizes:
        return default
    counts: dict[float, int] = {}
    for value in sizes:
        key = round(float(value) * 2) / 2
        counts[key] = counts.get(key, 0) + 1
    return max(counts.items(), key=lambda item: (item[1], item[0]))[0]


def _measure_size(scratch_page: Any, rect: Any, text: str, base: float) -> float:
    floor = max(4.0, base * 0.5)
    size = max(4.0, min(base, 24.0))
    while size >= floor:
        remaining = scratch_page.insert_textbox(
            rect, text, fontname="F0", fontsize=size, color=(0, 0, 0), align=0, render_mode=3
        )
        if remaining >= 0:
            return size
        size -= 0.25
    return floor


def _overlap_top(page_units: list[Unit], index: int, page_rect: Any) -> float:
    rect = fitz.Rect(*page_units[index].location["bbox"])
    top = page_rect.y1 - 36
    for other in page_units[index + 1 :]:
        other_rect = fitz.Rect(*other.location["bbox"])
        overlap = min(rect.x1, other_rect.x1) - max(rect.x0, other_rect.x0)
        if overlap > 0.3 * min(rect.width, other_rect.width):
            top = min(top, other_rect.y0 - 4)
            break
    return top


def overlay_pdf(parsed: ParsedDocument, translations: dict[str, str], out: Path) -> int:
    if fitz is None:
        raise DocumentError("parser_missing", "PyMuPDF is not installed")
    source = parsed.handle.get("source")
    if not source:
        raise DocumentError("parse_failed", "missing pdf source path")

    document = fitz.open(str(source))
    font = fitz.Font("cjk")

    max_width = max((page.rect.width for page in document), default=595.0)
    max_height = max((page.rect.height for page in document), default=842.0)
    scratch = fitz.open()
    scratch_page = scratch.new_page(width=max_width, height=max_height)
    scratch_page.insert_font(fontname="F0", fontbuffer=font.buffer)

    items: list[dict[str, Any]] = []
    for page_index in range(document.page_count):
        page = document[page_index]
        page_units = [unit for unit in parsed.units if int(unit.location.get("page", -1)) == page_index]
        page_units.sort(key=lambda unit: unit.location.get("bbox", [0, 0, 0, 0])[1])
        for index, unit in enumerate(page_units):
            translated = (translations.get(unit.id) or "").strip()
            if not translated or translated == unit.text.strip():
                continue
            bbox = unit.location.get("bbox")
            if not bbox:
                continue
            rect = fitz.Rect(*bbox)
            if rect.is_empty or rect.width < 1 or rect.height < 1:
                continue
            if rect.width >= 0.55 * page.rect.width:
                top = _overlap_top(page_units, index, page.rect)
                bottom = min(top, rect.y0 + rect.height * 2.5, page.rect.y1 - 36)
                if bottom > rect.y1 + 4:
                    rect = fitz.Rect(rect.x0, rect.y0, rect.x1, bottom)
            items.append(
                {
                    "page": page_index,
                    "rect": rect,
                    "text": translated,
                    "base": float(unit.location.get("font_size") or 10.0),
                    "color": int(unit.location.get("color") or 0),
                }
            )

    groups: dict[float, list[dict[str, Any]]] = {}
    for item in items:
        key = round(item["base"] * 2) / 2
        groups.setdefault(key, []).append(item)

    for group in groups.values():
        base = sorted(item["base"] for item in group)[len(group) // 2]
        measured = [_measure_size(scratch_page, item["rect"], item["text"], base) for item in group]
        cluster_size = min(measured)
        for item in group:
            item["size"] = cluster_size

    scratch.close()

    replaced = 0
    for page_index in range(document.page_count):
        page_items = [item for item in items if item["page"] == page_index]
        if not page_items:
            continue
        page = document[page_index]
        for item in page_items:
            page.add_redact_annot(item["rect"], fill=None)
        page.apply_redactions(
            images=fitz.PDF_REDACT_IMAGE_NONE,
            graphics=fitz.PDF_REDACT_LINE_ART_NONE,
        )
        page.insert_font(fontname="F0", fontbuffer=font.buffer)
        for item in page_items:
            color = _rgb(item["color"])
            try:
                remaining = page.insert_textbox(
                    item["rect"], item["text"], fontname="F0", fontsize=item["size"], color=color, align=0
                )
            except Exception:
                remaining = -1
            if remaining < 0:
                try:
                    page.insert_text(
                        (item["rect"].x0, item["rect"].y0 + item["size"]),
                        item["text"],
                        fontname="F0",
                        fontsize=item["size"],
                        color=color,
                    )
                except Exception:
                    pass
            replaced += 1

    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out), garbage=3, deflate=True)
    document.close()
    return replaced


class PdfAdapter:
    extensions = (".pdf",)

    def parse(self, path: Path, options: dict[str, Any] | None = None) -> ParsedDocument:
        if fitz is None:
            raise DocumentError("parser_missing", "PyMuPDF is not installed")
        try:
            document = fitz.open(str(path))
        except Exception as exc:
            raise DocumentError("parse_failed", "cannot open pdf: " + str(exc)) from exc
        if getattr(document, "needs_pass", False):
            document.close()
            raise DocumentError("encrypted_document", "pdf is password protected")

        units: list[Unit] = []
        pages_with_text = 0
        empty_pages = 0
        width = 595.0
        height = 842.0

        for page_index in range(document.page_count):
            page = document[page_index]
            if page_index == 0:
                width = float(page.rect.width)
                height = float(page.rect.height)
            blocks = [block for block in (page.get_text("dict").get("blocks") or []) if block.get("type") == 0]
            blocks.sort(key=lambda block: (round(block["bbox"][1], 1), round(block["bbox"][0], 1)))
            page_units = 0
            for block_index, block in enumerate(blocks):
                line_texts: list[str] = []
                sizes: list[float] = []
                colors: list[int] = []
                for line in block.get("lines") or []:
                    spans = line.get("spans") or []
                    line_text = "".join(span.get("text", "") for span in spans).strip()
                    if line_text:
                        line_texts.append(line_text)
                    for span in spans:
                        if span.get("size"):
                            sizes.append(float(span["size"]))
                        if span.get("color") is not None:
                            colors.append(int(span["color"]))
                text = flow_text("\n".join(line_texts))
                if len(text) < 2:
                    continue
                size = _dominant_size(sizes)
                color = colors[0] if colors else 0
                units.append(
                    Unit(
                        id="p" + str(page_index) + "b" + str(block_index),
                        text=text,
                        location={
                            "type": "pdf_block",
                            "page": page_index,
                            "block": block_index,
                            "bbox": [float(value) for value in block["bbox"]],
                            "font_size": float(size),
                            "color": color,
                        },
                    )
                )
                page_units += 1
            if page_units:
                pages_with_text += 1
            else:
                empty_pages += 1

        meta = {
            "pages": document.page_count,
            "pages_with_text": pages_with_text,
            "empty_pages": empty_pages,
            "width": width,
            "height": height,
        }
        document.close()
        return ParsedDocument(units=units, meta=meta, handle={"source": str(path)})

    def rebuild(self, parsed: ParsedDocument, translations: dict[str, str], out: Path) -> None:
        mode = os.environ.get("PDF_OUTPUT_MODE", "overlay").strip().lower()
        if mode == "reflow":
            temp_docx = out.with_name(out.stem + ".source.docx")
            try:
                build_docx(parsed, translations, temp_docx)
                convert_docx_to_pdf(temp_docx, out)
            finally:
                temp_docx.unlink(missing_ok=True)
            return
        overlay_pdf(parsed, translations, out)

    def verify(self, source: Path, out: Path, parsed: ParsedDocument) -> list[dict[str, Any]]:
        if fitz is None:
            return [{"check": "pdf_library", "ok": False, "detail": "PyMuPDF missing"}]
        try:
            document = fitz.open(str(out))
        except Exception as exc:
            return [{"check": "reopen", "ok": False, "detail": str(exc)}]
        text_chars = 0
        for page in document:
            text_chars += len(page.get_text().strip())
        pages = document.page_count
        document.close()
        expected_pages = int(parsed.meta.get("pages", 0))
        return [
            {"check": "page_count", "ok": pages == expected_pages, "source": expected_pages, "output": pages},
            {"check": "has_text", "ok": text_chars > 0, "chars": text_chars},
        ]
