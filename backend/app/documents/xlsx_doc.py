from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .base import DocumentError, ParsedDocument, Unit


class XlsxAdapter:
    extensions = (".xlsx", ".xlsm")

    def _load(self, path: Path):
        return load_workbook(filename=str(path), data_only=False, keep_vba=path.suffix.lower() == ".xlsm")

    def parse(self, path: Path, options: dict[str, Any] | None = None) -> ParsedDocument:
        try:
            workbook = self._load(path)
        except Exception as exc:
            raise DocumentError("parse_failed", "cannot open workbook: " + str(exc)) from exc

        units: list[Unit] = []
        formulas: set[str] = set()
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas.add(sheet.title + "!" + cell.coordinate)
                        continue
                    value = cell.value
                    if not isinstance(value, str) or not value.strip():
                        continue
                    units.append(
                        Unit(
                            id=sheet.title + "!" + cell.coordinate,
                            text=value,
                            location={"type": "cell", "sheet": sheet.title, "cell": cell.coordinate},
                        )
                    )

        meta = {"sheets": list(workbook.sheetnames), "formulas": len(formulas)}
        return ParsedDocument(units=units, meta=meta, handle={"workbook": workbook, "formulas": formulas})

    def rebuild(self, parsed: ParsedDocument, translations: dict[str, str], out: Path) -> None:
        workbook = parsed.handle["workbook"]
        for unit_id, translated in translations.items():
            if "!" not in unit_id:
                continue
            sheet_name, coord = unit_id.rsplit("!", 1)
            if sheet_name not in workbook.sheetnames:
                continue
            workbook[sheet_name][coord] = translated
        out.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(str(out))

    def verify(self, source: Path, out: Path, parsed: ParsedDocument) -> list[dict[str, Any]]:
        try:
            workbook = self._load(out)
        except Exception as exc:
            return [{"check": "reopen", "ok": False, "detail": str(exc)}]

        checks: list[dict[str, Any]] = []
        expected_sheets = list(parsed.meta.get("sheets") or [])
        checks.append(
            {
                "check": "sheets",
                "ok": list(workbook.sheetnames) == expected_sheets,
                "source": expected_sheets,
                "output": list(workbook.sheetnames),
            }
        )

        formulas: set[str] = set()
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formulas.add(sheet.title + "!" + cell.coordinate)
        checks.append(
            {
                "check": "formula_hash",
                "ok": formulas == parsed.handle["formulas"],
                "source_count": len(parsed.handle["formulas"]),
                "output_count": len(formulas),
            }
        )
        return checks
