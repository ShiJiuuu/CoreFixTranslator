from __future__ import annotations

import zipfile
from pathlib import Path
from typing import Any

from docx import Document

from .base import DocumentError, ParsedDocument, Unit


def _set_paragraph_text(paragraph, text: str) -> None:
    runs = paragraph.runs
    if runs:
        runs[0].text = text
        for run in runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


class DocxAdapter:
    extensions = (".docx",)

    def parse(self, path: Path, options: dict[str, Any] | None = None) -> ParsedDocument:
        options = options or {}
        try:
            doc = Document(str(path))
        except Exception as exc:
            raise DocumentError("parse_failed", "cannot open docx: " + str(exc)) from exc

        mapping: dict[str, Any] = {}
        units: list[Unit] = []
        skipped_links = 0

        for index, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            if getattr(paragraph, "hyperlinks", None):
                skipped_links += 1
                continue
            unit_id = "p" + str(index)
            units.append(Unit(id=unit_id, text=text, location={"type": "paragraph", "index": index}))
            mapping[unit_id] = paragraph

        for ti, table in enumerate(doc.tables):
            for ri, row in enumerate(table.rows):
                for ci, cell in enumerate(row.cells):
                    for pi, paragraph in enumerate(cell.paragraphs):
                        text = paragraph.text.strip()
                        if not text:
                            continue
                        unit_id = "t" + str(ti) + "r" + str(ri) + "c" + str(ci) + "p" + str(pi)
                        units.append(
                            Unit(
                                id=unit_id,
                                text=text,
                                location={"type": "table_cell", "table": ti, "row": ri, "col": ci},
                            )
                        )
                        mapping[unit_id] = paragraph

        headers = 0
        if options.get("translate_headers", False):
            for si, section in enumerate(doc.sections):
                for pi, paragraph in enumerate(section.header.paragraphs):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    unit_id = "h" + str(si) + "p" + str(pi)
                    units.append(Unit(id=unit_id, text=text, location={"type": "header", "section": si}))
                    mapping[unit_id] = paragraph
                    headers += 1
                for pi, paragraph in enumerate(section.footer.paragraphs):
                    text = paragraph.text.strip()
                    if not text:
                        continue
                    unit_id = "f" + str(si) + "p" + str(pi)
                    units.append(Unit(id=unit_id, text=text, location={"type": "footer", "section": si}))
                    mapping[unit_id] = paragraph
                    headers += 1

        meta = {
            "paragraphs": len(doc.paragraphs),
            "tables": len(doc.tables),
            "header_footer_units": headers,
            "skipped_hyperlinks": skipped_links,
        }
        return ParsedDocument(units=units, meta=meta, handle={"doc": doc, "mapping": mapping})

    def rebuild(self, parsed: ParsedDocument, translations: dict[str, str], out: Path) -> None:
        doc = parsed.handle["doc"]
        mapping = parsed.handle["mapping"]
        for unit_id, translated in translations.items():
            paragraph = mapping.get(unit_id)
            if paragraph is None:
                continue
            _set_paragraph_text(paragraph, translated)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(out))

    def verify(self, source: Path, out: Path, parsed: ParsedDocument) -> list[dict[str, Any]]:
        checks: list[dict[str, Any]] = []
        try:
            with zipfile.ZipFile(out) as archive:
                bad = archive.testzip()
            checks.append({"check": "zip_integrity", "ok": bad is None, "detail": bad})
        except Exception as exc:
            checks.append({"check": "zip_integrity", "ok": False, "detail": str(exc)})
            return checks

        try:
            doc = Document(str(out))
        except Exception as exc:
            checks.append({"check": "reopen", "ok": False, "detail": str(exc)})
            return checks

        expected_paragraphs = int(parsed.meta.get("paragraphs", 0))
        expected_tables = int(parsed.meta.get("tables", 0))
        checks.append(
            {
                "check": "paragraph_count",
                "ok": len(doc.paragraphs) == expected_paragraphs,
                "source": expected_paragraphs,
                "output": len(doc.paragraphs),
            }
        )
        checks.append(
            {
                "check": "table_count",
                "ok": len(doc.tables) == expected_tables,
                "source": expected_tables,
                "output": len(doc.tables),
            }
        )
        return checks
