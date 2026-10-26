from pathlib import Path

from docx import Document

from backend.app.config import env_defaults
from backend.app.model import FakeAdapter
from backend.app.pipeline import translate_document


def _make_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Report", level=1)
    document.add_paragraph("Revenue grew 12% in 2024.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    document.save(str(path))


def test_docx_roundtrip(tmp_path: Path):
    source = tmp_path / "in.docx"
    out = tmp_path / "out.docx"
    _make_docx(source)

    settings = env_defaults()
    report = translate_document(
        source,
        out,
        settings=settings,
        model=FakeAdapter(settings),
        source_lang="en",
        target_lang="zh-CN",
    )

    assert out.exists()
    reopened = Document(str(out))
    assert len(reopened.paragraphs) == 2
    assert len(reopened.tables) == 1
    assert reopened.paragraphs[0].text.startswith("[zh-CN]")
    assert reopened.tables[0].cell(0, 0).text.startswith("[zh-CN]")
    assert reopened.tables[0].cell(0, 1).text.startswith("[zh-CN]")
    assert report["hard_checks_ok"] is True
    assert report["units"] >= 4


def test_nested_table_cell_translated(tmp_path: Path):
    source = tmp_path / "nested.docx"
    out = tmp_path / "nested.out.docx"
    document = Document()
    outer = document.add_table(rows=1, cols=1)
    inner = outer.cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Nested cell"
    document.save(str(source))

    settings = env_defaults()
    report = translate_document(
        source,
        out,
        settings=settings,
        model=FakeAdapter(settings),
        source_lang="en",
        target_lang="zh-CN",
    )

    reopened = Document(str(out))
    nested = reopened.tables[0].cell(0, 0).tables[0]
    assert nested.cell(0, 0).text.startswith("[zh-CN]")
    assert report["units"] == 1
