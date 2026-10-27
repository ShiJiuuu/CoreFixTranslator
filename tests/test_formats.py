from pathlib import Path

import fitz

import pytest
from docx import Document

from backend.app.documents import output_suffix_for
from backend.app.documents.base import DocumentError, ParsedDocument, Unit
from backend.app.config import env_defaults
from backend.app.documents.pdf_doc import build_docx, convert_docx_to_pdf, flow_text
from backend.app.model.fake import FakeAdapter
from backend.app.pipeline import translate_document


def test_pdf_outputs_pdf():
    assert output_suffix_for(Path("report.pdf")) == ".pdf"


def test_docx_keeps_docx():
    assert output_suffix_for(Path("report.docx")) == ".docx"


def test_text_keeps_suffix():
    assert output_suffix_for(Path("notes.txt")) == ".txt"


def test_pdf_builds_flowing_docx(tmp_path: Path):
    parsed = ParsedDocument(
        units=[
            Unit(id="p0b0", text="Hello", location={"page": 0}),
            Unit(id="p1b0", text="World", location={"page": 1}),
        ],
        meta={"pages": 2, "width": 595.0, "height": 842.0},
    )
    out = tmp_path / "flow.docx"
    build_docx(parsed, {"p0b0": "第一段", "p1b0": "第二段"}, out)

    document = Document(str(out))
    assert [p.text for p in document.paragraphs] == ["第一段", "第二段"]
    assert round(document.sections[0].page_width.pt) == 595


def test_pdf_converter_missing(tmp_path: Path):
    source = tmp_path / "empty.docx"
    Document().save(str(source))
    with pytest.raises(DocumentError) as info:
        convert_docx_to_pdf(source, tmp_path / "empty.pdf", binary="/nonexistent/soffice")
    assert info.value.code == "pdf_converter_missing"


def test_flow_text_joins_lines():
    assert flow_text("Hello\nworld") == "Hello world"
    assert flow_text("第一行\n第二行") == "第一行第二行"
    assert flow_text("Line 1\n\nLine 2") == "Line 1 Line 2"


class _VaryingAdapter(FakeAdapter):
    def translate_batch(self, units, **kwargs):
        lengths = [6, 14, 26]
        result = {}
        for index, unit in enumerate(units):
            result[unit["id"]] = "译文内容" * lengths[index % len(lengths)]
        return result, {}


def test_pdf_overlay_uniform_sizes(tmp_path: Path):
    source = tmp_path / "in.pdf"
    document = fitz.open()
    page = document.new_page()
    for top in (90, 170, 250):
        page.insert_text((72, top), "Same style body text here.", fontsize=12)
    document.save(str(source))
    document.close()

    settings = env_defaults()
    out = tmp_path / "out.pdf"
    translate_document(
        source,
        out,
        settings=settings,
        model=_VaryingAdapter(settings),
        source_lang="en",
        target_lang="zh-CN",
    )

    check = fitz.open(str(out))
    sizes = set()
    for block in check[0].get_text("dict")["blocks"]:
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                sizes.add(round(span["size"], 2))
    check.close()
    assert len(sizes) == 1
