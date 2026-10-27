from pathlib import Path

import fitz
from docx import Document

from backend.app.config import env_defaults
from backend.app.documents.pdf.parser import parse_pdf
from backend.app.documents.pdf.render_docx import render_docx
from backend.app.documents.pdf.text import flow_text
from backend.app.model import FakeAdapter
from backend.app.pipeline import translate_document


def _make_pdf(path: Path) -> None:
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text((60, 80), "Document Title", fontsize=20)
    page.insert_text((60, 120), "Body paragraph one.", fontsize=10.5)
    page.insert_text((60, 150), "Body paragraph two.", fontsize=10.5)
    for index, text in enumerate(["Left column one", "Left column two"]):
        page.insert_text((60, 220 + index * 24), text, fontsize=10.5)
    for index, text in enumerate(["Right column one", "Right column two"]):
        page.insert_text((320, 220 + index * 24), text, fontsize=10.5)
    document.save(str(path))
    document.close()


def test_flow_text_joins_lines():
    assert flow_text("Hello\nworld") == "Hello world"
    assert flow_text("第一行\n第二行") == "第一行第二行"


def test_heading_and_body_classification(tmp_path: Path):
    source = tmp_path / "doc.pdf"
    _make_pdf(source)

    ir = parse_pdf(source)
    blocks = ir.ordered_blocks()
    assert blocks[0].type == "heading"
    assert blocks[0].level == 1
    assert any(block.type == "paragraph" for block in blocks)


def test_two_column_reading_order(tmp_path: Path):
    source = tmp_path / "columns.pdf"
    _make_pdf(source)

    ir = parse_pdf(source)
    texts = [block.source_text for block in ir.ordered_blocks()]
    assert texts.index("Left column one") < texts.index("Right column one")
    assert texts.index("Left column two") < texts.index("Right column two")


def test_render_docx_from_ir(tmp_path: Path):
    source = tmp_path / "doc.pdf"
    _make_pdf(source)
    ir = parse_pdf(source)
    for block in ir.ordered_blocks():
        block.target_text = "[zh] " + block.source_text

    out = tmp_path / "out.docx"
    render_docx(ir, out)

    document = Document(str(out))
    assert document.paragraphs[0].style.name == "Heading 1"
    assert document.paragraphs[0].text == "[zh] Document Title"
    assert len(document.paragraphs) == len(ir.ordered_blocks())


def test_pdf_pipeline_outputs_docx(tmp_path: Path, monkeypatch):
    monkeypatch.setenv("PDF_RENDER", "docx")
    source = tmp_path / "doc.pdf"
    _make_pdf(source)
    out = tmp_path / "out.docx"

    settings = env_defaults()
    report = translate_document(
        source,
        out,
        settings=settings,
        model=FakeAdapter(settings),
        source_lang="en",
        target_lang="zh-CN",
    )

    assert out.exists()
    document = Document(str(out))
    assert any("[zh-CN] Document Title" in paragraph.text for paragraph in document.paragraphs)
    assert report["hard_checks_ok"] is True
    assert report["units"] >= 5


def test_list_items_split_into_blocks(tmp_path: Path):
    source = tmp_path / "list.pdf"
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    page.insert_text(
        (60, 120),
        "• First bullet item\n• Second bullet item\n• Third bullet item",
        fontsize=12,
    )
    document.save(str(source))
    document.close()

    ir = parse_pdf(source)
    items = [block for block in ir.ordered_blocks() if block.type == "listItem"]
    assert len(items) == 3
    assert items[0].source_text.endswith("First bullet item")
    assert items[2].source_text.endswith("Third bullet item")


def test_sparse_line_spacing_stays_one_block(tmp_path: Path):
    source = tmp_path / "spacing.pdf"
    document = fitz.open()
    page = document.new_page(width=595, height=842)
    for index, text in enumerate(
        ["First line of paragraph", "Second line continues", "Third line ends"]
    ):
        page.insert_text((60, 120 + index * 18), text, fontsize=12)
    document.save(str(source))
    document.close()

    ir = parse_pdf(source)
    blocks = ir.ordered_blocks()
    assert len(blocks) == 1
    assert "First line of paragraph" in blocks[0].source_text
    assert "Third line ends" in blocks[0].source_text
